from __future__ import annotations

from email.message import EmailMessage
from pathlib import Path

from odf.opendocument import OpenDocumentText
from odf.text import P
from pptx import Presentation

from tooru.cloud.intelligence import extract_document


def _text(chunks) -> str:
    return "\n".join(chunk.text for chunk in chunks)


def test_extracts_pptx(tmp_path: Path) -> None:
    path = tmp_path / "presentation.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )
    slide.shapes.title.text = "Dragon Tory presentation"
    textbox = slide.shapes.add_textbox(0, 0, 100, 100)
    textbox.text = "Contract deadline 2026-12-01"
    presentation.save(path)

    chunks = extract_document(
        path,
        name=path.name,
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "presentationml.presentation"
        ),
    )

    text = _text(chunks)
    assert "Dragon Tory presentation" in text
    assert "2026-12-01" in text


def test_extracts_rtf_html_odt_and_eml(tmp_path: Path) -> None:
    rtf = tmp_path / "note.rtf"
    rtf.write_text(
        r"{\rtf1\ansi Dragon Tory RTF document}",
        encoding="ascii",
    )
    assert "Dragon Tory RTF" in _text(
        extract_document(
            rtf,
            name=rtf.name,
            content_type="application/rtf",
        )
    )

    html = tmp_path / "page.html"
    html.write_text(
        "<html><body><h1>Тоору</h1><p>Документ HTML</p></body></html>",
        encoding="utf-8",
    )
    html_text = _text(
        extract_document(
            html,
            name=html.name,
            content_type="text/html",
        )
    )
    assert "Тоору" in html_text
    assert "Документ HTML" in html_text

    odt = tmp_path / "document.odt"
    document = OpenDocumentText()
    document.text.addElement(P(text="Dragon Tory OpenDocument"))
    document.save(str(odt))
    assert "Dragon Tory OpenDocument" in _text(
        extract_document(
            odt,
            name=odt.name,
            content_type="application/vnd.oasis.opendocument.text",
        )
    )

    eml = tmp_path / "mail.eml"
    message = EmailMessage()
    message["Subject"] = "Договор поставки"
    message["From"] = "supplier@example.com"
    message["To"] = "user@example.com"
    message.set_content("Срок поставки 15 декабря 2026 года.")
    eml.write_bytes(message.as_bytes())
    eml_text = _text(
        extract_document(
            eml,
            name=eml.name,
            content_type="message/rfc822",
        )
    )
    assert "Договор поставки" in eml_text
    assert "15 декабря 2026" in eml_text
