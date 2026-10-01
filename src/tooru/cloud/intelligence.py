from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


@dataclass(slots=True)
class ExtractedChunk:
    label: str
    text: str
    page: int | None = None


class UnsupportedDocumentError(ValueError):
    pass


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _split_text(
    text: str,
    *,
    label: str,
    page: int | None = None,
    max_chars: int = 3_500,
) -> list[ExtractedChunk]:
    cleaned = "\n".join(
        line.rstrip()
        for line in text.replace("\x00", "").splitlines()
    ).strip()
    if not cleaned:
        return []

    chunks: list[ExtractedChunk] = []
    start = 0
    part = 1
    while start < len(cleaned):
        end = min(len(cleaned), start + max_chars)
        if end < len(cleaned):
            boundary = cleaned.rfind("\n", start, end)
            if boundary <= start + max_chars // 2:
                boundary = cleaned.rfind(" ", start, end)
            if boundary > start:
                end = boundary
        piece = cleaned[start:end].strip()
        if piece:
            suffix = f" · часть {part}" if len(cleaned) > max_chars else ""
            chunks.append(
                ExtractedChunk(
                    label=label + suffix,
                    text=piece,
                    page=page,
                )
            )
            part += 1
        start = max(end, start + 1)
    return chunks


def extract_document(
    path: Path,
    *,
    name: str,
    content_type: str,
) -> list[ExtractedChunk]:
    suffix = Path(name).suffix.lower()
    if suffix in {".txt", ".md", ".json", ".log", ".yaml", ".yml"}:
        return _split_text(
            _decode_text(path.read_bytes()),
            label="текст",
        )

    if suffix == ".csv":
        raw = _decode_text(path.read_bytes())
        reader = csv.reader(io.StringIO(raw))
        text = "\n".join(" | ".join(row) for row in reader)
        return _split_text(text, label="таблица CSV")

    if suffix == ".pdf" or content_type == "application/pdf":
        reader = PdfReader(str(path))
        chunks: list[ExtractedChunk] = []
        for number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            chunks.extend(
                _split_text(
                    text,
                    label=f"страница {number}",
                    page=number,
                )
            )
        return chunks

    if suffix == ".docx":
        document = Document(str(path))
        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(
                    " | ".join(cell.text for cell in row.cells)
                )
        return _split_text(
            "\n".join(paragraphs),
            label="документ DOCX",
        )

    if suffix == ".xlsx":
        workbook = load_workbook(
            filename=str(path),
            read_only=True,
            data_only=True,
        )
        chunks: list[ExtractedChunk] = []
        try:
            for sheet in workbook.worksheets:
                rows: list[str] = []
                for row in sheet.iter_rows(values_only=True):
                    values = [
                        "" if value is None else str(value)
                        for value in row
                    ]
                    if any(values):
                        rows.append(" | ".join(values))
                chunks.extend(
                    _split_text(
                        "\n".join(rows),
                        label=f"лист {sheet.title}",
                    )
                )
        finally:
            workbook.close()
        return chunks

    raise UnsupportedDocumentError(
        "Этот формат пока нельзя индексировать. "
        "Поддерживаются TXT, MD, JSON, CSV, PDF, DOCX и XLSX."
    )


def preview_document(
    path: Path,
    *,
    name: str,
    content_type: str,
    max_chars: int = 30_000,
) -> dict:
    chunks = extract_document(
        path,
        name=name,
        content_type=content_type,
    )
    used: list[dict] = []
    total = 0
    for chunk in chunks:
        remaining = max_chars - total
        if remaining <= 0:
            break
        text = chunk.text[:remaining]
        used.append(
            {
                "label": chunk.label,
                "page": chunk.page,
                "text": text,
            }
        )
        total += len(text)
    return {
        "chunks": used,
        "chunk_count": len(chunks),
        "preview_chars": total,
        "truncated": len(used) < len(chunks)
        or sum(len(c.text) for c in chunks) > total,
    }
